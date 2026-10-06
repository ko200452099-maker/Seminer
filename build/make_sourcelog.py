#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Builds 04_Source_Log.docx — every number on the slides, traced to its source with the exact quote."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x10,0x2A,0x43); TEAL = RGBColor(0x1F,0x7A,0x84)
GRAY = RGBColor(0x5A,0x66,0x72); RED = RGBColor(0xB4,0x3B,0x2E); AMBER = RGBColor(0xA8,0x6B,0x00)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(9)
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.0
for name, size, color in (("Heading 1", 13, NAVY), ("Heading 2", 11, TEAL)):
    sty = doc.styles[name]; sty.font.name = "Calibri"; sty.font.size = Pt(size); sty.font.bold = True
    sty.font.color.rgb = color; sty.paragraph_format.space_before = Pt(8); sty.paragraph_format.space_after = Pt(3)
sec = doc.sections[0]
sec.top_margin = Cm(1.4); sec.bottom_margin = Cm(1.4); sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)

def para(text, size=9, bold=False, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    if color: r.font.color.rgb = color
    return p

def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexc); tcPr.append(shd)

def tbl(rows, widths, fs=8):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci); cell.text = ""
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(val)); r.font.size = Pt(fs); r.font.name = "Calibri"
            if ri == 0:
                r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); shade(cell, "102A43")
            elif ri % 2 == 0:
                shade(cell, "F4F7F9")
    total = sum(widths)
    for ci, w in enumerate(widths):
        for ri in range(len(rows)):
            t.cell(ri, ci).width = Cm(18.0 * w / total)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

# ---------------- header ----------------
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SOURCE LOG — EVERY NUMBER ON THE SLIDES, TRACED TO ITS PAPER"); r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = TEAL
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Asphaltenes: Predicting and Preventing Wellbore Plugging"); r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = NAVY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Print this and keep it at the lectern. If an examiner challenges any figure, hand them this page — not a guess.")
r.font.size = Pt(9); r.font.italic = True; r.font.color.rgb = GRAY
para("Status key:   [V] = quote verified verbatim from an accessible full text or publisher abstract during this build.   [L] = full-text / page-number confirmation recommended at the university library in Week 1 (marked on the item). "
     "UPDATED after the technical cross-check: the de Boer definition and the ASTM D6703 entry were CORRECTED (rows marked above); the D6560/IP 143 editions are confirmed. The full verification record is 06_Technical_Accuracy_and_Risk_Register.docx, Part A.",
     size=8, italic=True, color=AMBER, space_after=6)

rows = [["Slide", "Claim on the slide", "Source", "Exact wording (verbatim)", "Status"]]

data = [
 ("2, 17", "“billions of dollars every year” — industry-wide cost", "Farooq, U.; Lædre, S.; Gawel, K. Review of Asphaltenes in an Electric Field. Energy & Fuels 2021, 35(9), 7285–7304. doi:10.1021/acs.energyfuels.0c03962",
  "“Asphaltene deposition costs the oil industry billions of dollars every year, including reduction of production capacity, wells shut-in, and implications of the management techniques.”", "[V]"),
 ("2, 17", "≈ $70 M: one deepwater GoM well, shut-in + cleanup", "Farooq, U.; Lædre, S.; Gawel, K. (2021), Energy & Fuels, 35(9), 7285–7304",
  "“it was reported from an oil field in the Gulf of Mexico that the cost of a well shut-in for a cleanup operation due to asphaltene deposition was approximately $70 million/well.”", "[V]"),
 ("17", "up to $100 M if deposit forms at the SCSSV; removal $0.3–3.5 M/well; sidetrack ≈ $50 M; downtime ≈ $700 k/day", "Asphaltene Deposition Control by Chemical Inhibitors, 1st ed.; Gulf Professional/Elsevier, 2021; Ch.1 (economic analysis)",
  "“the rig intervention costs were estimated to be equal to $70 million per well… could even increase to $100 million per well if the deposits are formed at the surface-controlled subsurface safety valve.” Table: removal $300,000–3,500,000/well; sidetrack $50,000,000/well; downtime $700,000/day (7,000 BPD).", "[V]"),
 ("2, 17", "up to ≈ $1.2 M/day deferred production", "Stratiev, D. et al. Mitigation of Asphaltene Deposit Formation via Chemical Additives: A Review. Processes 2025, 13(1), 141. doi:10.3390/pr13010141",
  "“the financial loss due to lost production can be as high as USD 1.2 M/day, even with the scenario of a USD 30/bbl oil price.”", "[V]"),
 ("17", "treatment documented up to ≈ $3 M/well", "Stratiev, D. et al. (2025), Processes, 13(1), 141",
  "“the cost of treatment of asphaltene deposits in the well can be as high as USD 3 M.”", "[V]"),
 ("17", "inhibitor: $31–46 k/well/yr (Middle East); $330–390 k/well/yr (GoM) — attribution conflict", "Khaleel, A.T.; Abutaqiya, M.I.L.; Sisco, C.J.; Vargas, F.M. Mitigation of Asphaltene Deposition by Re-injection of Dead Oil. Fluid Phase Equilibria 2020, 514, 112552. doi:10.1016/j.fluid.2020.112552 (citing Cenegy, 2001)",
  "“These chemicals are also expensive, costing between $31,000 – $46,000 per well per year for fields in the Middle East and $330,000 – $390,000 per well per year for fields in the Gulf of Mexico.” CONFLICT: Farooq et al. (2021) reverse the geography. The Elsevier monograph (2021) supports the majority attribution.  [L] confirm Cenegy’s original wording + SPE number.", "[V] [L]"),
 ("2, 4, 15", "20–25 % wellhead pressure lost in 15–20 days", "Haskett, C.E.; Tartera, M. SPE-994-PA / J. Pet. Technol. 1965, 17(4), 387–391. doi:10.2118/994-PA",
  "“Wells often lost 20 to 25 per cent of the wellhead pressure in 15 to 20 days, causing considerable loss in production.”  [L] page number from full text.", "[V] [L]"),
 ("15", "deposit composition: 83.4 % asphaltenes, 3.3 % carbenes, 13.3 % resins", "Haskett & Tartera (1965), SPE-994-PA",
  "“These deposits consisted of 83.4 per cent asphaltenes, 3.3 per cent carbenes, and 13.3 per cent resins and heavy fractions of oil, and were a plastic, tacky consistency.”", "[V]"),
 ("15", "reactive cleanup described as “expensive and cumbersome”", "Haskett & Tartera (1965), SPE-994-PA",
  "“This method, while satisfactory for cleaning the tubings, was expensive and cumbersome.”", "[V]"),
 ("2", "up to ⅔ of tubing radius: measured deposits, Hassi Messaoud", "Haskett & Tartera (1965) measurements, as restated in later SPE field-case literature",
  "“the deposit in five different wells of the Hassi Messaoud field was measured, and found to reach approximately two-thirds of the inner-tubing radius (Haskett and Tartera 1965).”  [L] cite the restating paper you actually consult.", "[V] [L]"),
 ("4, 13, 16", "GoM geometry: perfs 17,700–18,000 ft; mandrel 14,600 ft; 3,400 ft unprotected", "Wylde, J.; Punase, A. Asphaltenes: A Complex and Challenging Flow Assurance Issue To Measure and Quantify Risk. J. Pet. Technol., 30 Apr 2020 (Clariant Oil Services)",
  "“The well was perforated between 17,700 and 18,000 ft, while the downhole chemical injection mandrel was situated at 14,600 ft, resulting in 3,400 ft of unprotected tubing.”", "[V]"),
 ("4, 13", "tubing ID 2.30 → 1.35 in (≈ ⅔ of flow area lost)", "Wylde & Punase (2020), from their transient multiphase flow model",
  "“…which indicated reduction of the production tubing internal diameter from 2.3 to 1.35 in.”", "[V]"),
 ("13, 16", "dosage 750 → ≈ 500 ppm; +16 % production over 4 months", "Wylde & Punase (2020)",
  "“a solvent soak for 2 days… continuous AI-1 injection at a high dose rate of 750 ppm”; “a decrease in the readings to around 21 at 500-ppm dosage”; “16% overall increase in production (including overall shut-in time) over a 4-month period.”", "[V]"),
 ("3, 7", "CII = (Saturates + Asphaltenes) / (Resins + Aromatics); ≥ 0.9 unstable, < 0.7 stable", "Yen, A.; Yin, Y.R.; Asomaning, S. SPE 65376 (2001); thresholds after Asomaning (2003) — confirmed independently in three sources",
  "“If oil has a CII value below 0.7, it is defined as stable and if the CII is higher than 0.9 is considered as unstable.” Thresholds remain laboratory-specific.", "[V]"),
 ("10, 24", "UAOP sits above the bubble point; instability windows of thousands of psi (2,600–3,800 psi case)", "Fuel 2021, “Critical analysis of different techniques used to screen asphaltene stability in crude oils” [author list to confirm] + flow-assurance reference text",
  "“asphaltenes start to precipitate out from the crude oil at a pressure above the bubble point pressure and this pressure is termed as Upper Asphaltene Onset Pressure (UAOP).” / “some oils have an instability window of several thousand pounds per square inch… from 2600 to 3800 psi.”  [L] author list + pages.", "[V] [L]"),
 ("10, 24", "CORRECTED in the pack: the de Boer plot is a cross-plot of in-situ oil density (x) vs gas under-saturation, P_res − Pb (y), with severe / moderate / minimal risk zones", "de Boer, R.B.; Leerlooyer, K.; Eigner, M.R.P.; van Bergen, A.R.D. SPE Prod. Facil. 1995, 10(1), 55–61. doi:10.2118/24987-PA (restated identically in the J. Pet. Explor. Prod. Technol. 2018 onset-methods review and in J. Pet. Sci. Eng. 2022 SARA/CII correlation work)",
  "“It essentially evaluates the loss of asphaltene solubility as a reservoir fluid sample is depressurized. This PVT screen is a cross-plot of in situ density and the degree of under-saturation with respect to gas (the difference between reservoir and saturation pressures).” Paper conclusion: risk is greatest for “light crudes that are undersaturated with gas”. The onset-versus-bubble-point rule is a SEPARATE envelope screen.  [L] only the page number remains open.", "[V] [L]"),
 ("18", "“laboratory methods… are inadequate” critique of inhibitor testing", "Khaleel, A.T. et al. (2020), Fluid Phase Equilibria, 514, 112552",
  "“the laboratory methods used to validate the effectiveness of these chemicals are inadequate and the chemicals are often ineffective or actually worsen the problem when applied in the field.”", "[V]"),
 ("3, 12", "definition: insoluble in n-heptane, soluble in toluene (ASTM D6560 / IP 143)", "Wylde & Punase (2020) + ASTM D6560 / IP 143 scope",
  "“Asphaltenes represent a highly polar component of crude oil, insoluble in n-alkanes and soluble in aromatic solvents.” ASTM D6560-22 / IP 143:21 = Determination of Asphaltenes (Heptane Insolubles); D2007 = clay-gel SARA; D4124 = four fractions; D3279 = n-heptane insolubles; D6703-19 = Automated Heithaus Titrimetry — CORRECTED: D6703 is not a general precipitation test. The standard itself records that benzene was dropped as the aromatic solvent for health reasons.", "[V]"),
 ("3", "Nano-aggregates ≈ 2 nm, clusters ≈ 5 nm, molecular weight ≈ 750 g/mol (Yen–Mullins)", "Mullins, O.C. Energy & Fuels 2010, 24(4), 2179–2207; reviewed in Mullins et al. Energy & Fuels 2012, doi:10.1021/ef300185p",
  "“The most probable asphaltene molecular weight is ∼750 g/mol, with the island molecular architecture dominant… asphaltene molecules form nanoaggregates with an aggregation number less than 10. At higher concentrations, nanoaggregates form clusters…” Island-vs-archipelago remains an open debate — wording in the pack stays neutral.", "[V]"),
 ("15", "Thermal methods are not a melting route; ~497 °C steam questioned in the literature", "Mohammed, I.; Mahmoud, M.; Al Shehri, D.; El-Husseini, A.; Alade, O. J. Pet. Sci. Eng. 2021, 197, 107956, doi:10.1016/j.petrol.2020.107956",
  "“begs the question of the state of asphaltene as it does not melt but decompose into solid-like material which can pose problems within the formation”; “…the thermal methods are less efficient.” Counter-example anticipated in Q30.", "[V]"),
 ("15, 21", "BTEX HSE exposure flagged as a workplace-safety topic", "ACS Omega 2023, doi:10.1021/acsomega.3c03149; xylene clean-up field review",
  "“BTX chemicals … are acutely toxic and harmful to the environment”; “BTX solvents have low flash point, high acute toxicity and low biodegradability”; xylene clean-up “has limited effectiveness in addition to undesirable HSE effects.”", "[V]"),
 ("16", "Inhibitor vs dispersant: onset-shift vs aggregate-size reduction; neither dissolves a deposit", "Kelland, M.A. Production Chemicals for the Oil and Gas Industry, CRC Press 2009 (via AADE-24-FTCE-073)",
  "“Inhibitors typically affect the flocculation onset point, whereas dispersants usually reduce the size of aggregates formed. Inhibitors can also act as dispersants, but in general the inverse is not true.”", "[V]"),
 ("16", "Over-dosing can increase deposition — dose window, not dose maximum", "J. Pet. Sci. Eng. 2019, doi:10.1016/j.petrol.2019.106484",
  "“at high heptane vol%, inhibitor over-treatment (above recommended dose) increases deposition”; inhibitor self-aggregation reduces adsorption efficiency at excessive concentration.", "[V]"),
 ("20, 25", "NeqSim open-source CPA-EoS asphaltene screening (project fallback)", "NeqSim — Equinor, Apache-2.0 licence, v3.22. github.com/equinor/neqsim (pip install neqsim)",
  "“NeqSim — Non-Equilibrium Simulator. An open-source Java toolkit for thermodynamic calculations and process simulation.” Asphaltene stability via CPA EoS.", "[V]"),
]
for row in data:
    rows.append([str(x) for x in row])

tbl(rows, widths=[1.0, 3.4, 4.0, 8.6, 1.0], fs=7.5)

doc.add_heading("Week-1 library checklist (close every [L] before the first rehearsal)", level=1)
for t in [
 "Download SPE-994-PA (Haskett & Tartera, 1965) from OnePetro; log page numbers for the two verbatim quotes (20–25 % WHP loss; “expensive and cumbersome”) and the deposit composition.",
 "Confirm the Fuel (2021) “Critical analysis of different techniques…” author list, volume/page, and DOI.",
 "ASTM editions — CLOSED for citation: D6560-22 / IP 143:21, D2007, D4124, D3279, D6703-19 confirmed. Only action left: note which editions your library physically holds.",
 "Verify Cenegy (2001) SPE number and wording on OnePetro — it is the origin of the inhibitor-cost ranges and the source of the attribution conflict flagged above.",
 "Confirm the department’s software licences (PVTsim / Multiflash / OLGA-or-PIPESIM) and NeqSim install (pip install neqsim) — this is action item, not a citation.",
]:
    pp = doc.add_paragraph(style="List Bullet"); pp.paragraph_format.space_after = Pt(2)
    rr = pp.add_run(t); rr.font.size = Pt(8.5)

doc.add_heading("How to use this sheet in the defence", level=1)
para("Rule: never defend a number from memory alone — point to the line in this table. If a figure is challenged beyond the quote (e.g. “but which field?”), the honest answer is: “the review summarises it from Cenegy’s 2001 worldwide survey; the exact original wording is my Week-1 library task — I will send it to you.” That sentence is stronger than an improvised defence and it is exactly what a good researcher says.", size=9)

doc.save("/home/user/Seminar_Asphaltenes/04_Source_Log.docx")
print("saved source log")
